import os
from flask import render_template, request, redirect, url_for, Blueprint, current_app, make_response
from docx import Document

bp = Blueprint('main', __name__)

#Rota onde mostra os ficheiros, resumos individuais e resumo cronológico (se existirem)
@bp.route('/')
def index():
    manager = current_app.config['MANAGER']
    #Irá procurar a lista de ficheiros válidos
    ficheiros = manager.show_valid_files()
    #Os resumos individuais provenientes do Cop_LLM
    respostas = current_app.config.get('RESPOSTAS', None)
    #Os resumo final proveniente do Cop_Chronologic
    chron = current_app.config.get('CHRON', None)

    return render_template('index.html', files=ficheiros, respostas=respostas, chron=chron)

#Upload de ficheiros pelo utilizador utilizando o método POST de um formulário
@bp.route('/upload_files', methods=['POST'])
def upload_file():
    manager = current_app.config['MANAGER']
    #Irá receber a lista de ficheiros inseridos pelo utilizador
    files = request.files.getlist('files[]')
    
    for file in files:
        #E, para cada um dos ficheiros irá verificar a sua validação
        if file and manager.validator.is_valid(file.filename):
            #E se os mesmos forem válidos, serão guardados numa pasta uploads
            path = os.path.join(manager.upload_folder, file.filename)
            file.save(path)
            
    return redirect(url_for("main.index"))

#Gera os resumos individuais (um por ficheiro) usando LLM
@bp.route('/generate_results_llm', methods=['POST'])
def generate_results_llm():  
    manager = current_app.config['MANAGER']
    llm_manager = current_app.config['LLM_MANAGER']

    ficheiros = manager.show_valid_files()
    #Para cada um dos ficheiros inseridos na aplicação
    #os mesmos irão ser lidos
    conteudos = {f: manager.read_content(f) for f in ficheiros}

    #De seguida, com a ajuda do manager de llms, é feito, logo
    #após carregar no botão de gerar resposta, o processamento
    #do modelo CopLLM que irá resumir cada um dos ficheiros
    respostas = llm_manager.process_files_llm(conteudos)
    
    #Guarda os resumos num dicionario
    current_app.config['RESPOSTAS'] = respostas

    return redirect(url_for('main.index'))

#Gera o resumo cronológico final a partir dos resumos individuais
@bp.route('/generate_final_report', methods=['POST'])
def generate_final_report():
    llm_manager = current_app.config['LLM_MANAGER']
    respostas = current_app.config.get('RESPOSTAS', {})

    #Depois de se clicar no botão de gerar um relatório final,
    #o processamento vai buscar o dicionario de respostas anteriormente
    #gerado e envia tudo para o modelo CopChronologic para depois
    #se conseguir juntar tudo num documento e ordená-lo de forma
    #cronológica
    resumo_cronologico = llm_manager.process_files_llm_chron(respostas)

    #Guarda o resumo num dicionario
    current_app.config['CHRON'] = resumo_cronologico

    return render_template('relatorio_final.html', resposta=resumo_cronologico)

#Visualiza o conteúdo de um ficheiro
@bp.route('/ficheiro/<filename>')
def ver_ficheiro(filename):
    manager = current_app.config['MANAGER']
    #Lê-se o conteúdo original de um ficheiro já formatado
    conteudo = manager.read_content(filename)
    #E redireciona para uma página nova
    return render_template('ficheiro.html', filename=filename, conteudo=conteudo)

#Visualiza o conteúdo de uma resposta gerada de um ficheiro
@bp.route('/resposta/<filename>')
def ver_resposta(filename):
    respostas = current_app.config.get('RESPOSTAS', {})
    #Tendo em conta, o ícone em que o utilizador toca, será redirecionado
    #para outra página e nessa o resumo gerado será apresentado
    return render_template('resposta.html', filename=filename, resposta=respostas[filename])

#Edita manualmente uma resposta individual
@bp.route('/resposta/<filename>', methods=['POST'])
def editar_resposta(filename):
    respostas = current_app.config.get('RESPOSTAS', {})
    #O utilizador pode então editar a sua resposta, sendo 
    #criada uma nova variável com esse intuito
    nova_resposta = request.form.get('nova_resposta')

    #Se existir qualquer tipo de edição ao documento gerado,
    #essa resposta gerada irá ser substituida pela resposta
    #editada para depois ser considerada quando se gera
    #um relatório final
    if nova_resposta is not None:
        respostas[filename] = nova_resposta
        current_app.config['RESPOSTAS'] = respostas

    return redirect(url_for('main.ver_resposta', filename=filename))

#Edita manualmente um relatório final
@bp.route('/editar_chron', methods=['POST'])
def editar_chron():
    #O utilizador pode então editar o resumo final, sendo 
    #criada uma nova variável com esse intuito
    nova_resposta = request.form.get('nova_resposta')

    #Se existir qualquer tipo de edição ao documento gerado,
    #esse resumo final gerado irá ser substituido pelo resumo
    #editado para depois ser considerada quando se faz o download
    #do mesmo
    if nova_resposta is not None:
        current_app.config['CHRON'] = nova_resposta
    
    return redirect(url_for('main.ver_chron'))

#Visualiza o resumo cronológico final
@bp.route('/ver_chron')
def ver_chron():
    #A possibilidade de se visualizar o resumo final gerado
    resposta = current_app.config.get('CHRON', '')
    return render_template('relatorio_final.html', resposta=resposta, filename='relatorio_final')

#Faz download do resumo cronológico em formato .docx
@bp.route('/download_word')
def download_word():
    #Vai-se buscar o conteúdo que estaria guardado na aplicação
    resposta = current_app.config.get('CHRON', None)

    #A paste de testes para onde os downloads serão enviados
    pasta_relatorios = os.path.abspath('uploads/relatorios_finais')
    if not os.path.exists(pasta_relatorios):
        os.makedirs(pasta_relatorios)

    #E define o caminho completo para o ficheiro que será criado, com o seu nome
    caminho_ficheiro = os.path.join(pasta_relatorios, 'relatorio_final.docx')

    #Cria-se um novo documento Word com título e o conteúdo da resposta
    doc = Document()
    doc.add_heading('Relatório Final', level=1)
    doc.add_paragraph(resposta)
    #E salva-se então o ficheiro na pasta
    doc.save(caminho_ficheiro)

    #Abre o ficheiro criado em modo leitura binária
    with open(caminho_ficheiro, 'rb') as f:
        conteudo = f.read()

    #Prepara a resposta HTTP para que o browser faça o download do ficheiro
    response = make_response(conteudo)
    #Titulo do ficheiro
    response.headers.set('Content-Disposition', 'attachment', filename='relatorio_final.docx')
    #Tipo de ficheiro em Word, no caso
    response.headers.set('Content-Type', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    return response

# Apaga um ficheiro existente
@bp.route('/apagar_ficheiro/<filename>', methods=['POST'])
def apagar_ficheiro(filename):
    manager = current_app.config['MANAGER']
    path = os.path.join(manager.upload_folder, filename)
    if os.path.exists(path):
        os.remove(path)

    # Limpa também o conteúdo carregado na app se necessário
    respostas = current_app.config.get('RESPOSTAS', {})
    if filename in respostas:
        del respostas[filename]
        current_app.config['RESPOSTAS'] = respostas

    return redirect(url_for('main.index'))

#Ignora favicon requests
@bp.route('/favicon.ico')
def favicon():
    return '', 404